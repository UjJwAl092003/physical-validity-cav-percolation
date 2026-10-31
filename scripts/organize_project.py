import os
import shutil
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def main():
    print("Starting project decluttering and organization...")
    
    # 1. Create target directories
    for folder in ['figures', 'results', 'latex', 'docs']:
        os.makedirs(folder, exist_ok=True)
        print(f"Directory verified: {folder}/")

    # 2. Organize PNG images into figures/ (and keep copies in latex/ for local compilation)
    png_files = [
        'acceleration_and_curvature_distributions.png',
        'ir_comparison_across_conditions.png',
        'percolation_sweep_curve.png',
        'sample_trajectory_comparison.png'
    ]
    for png in png_files:
        if os.path.exists(png):
            dest_fig = os.path.join('figures', png)
            dest_latex = os.path.join('latex', png)
            shutil.copy2(png, dest_latex)
            shutil.move(png, dest_fig)
            print(f"Moved {png} -> figures/ (copy kept in latex/)")

    # 3. Organize Results (CSV and XLSX) into results/
    result_files = [
        'objective_1_results.xlsx',
        'objective_2_results.xlsx',
        'objective_1_summary_table.csv',
        'objective_1_all_seeds_detailed.csv',
        'objective_2_percolation_summary.csv',
        'objective_2_percolation_all_seeds.csv'
    ]
    for rf in result_files:
        if os.path.exists(rf):
            dest = os.path.join('results', rf)
            shutil.move(rf, dest)
            print(f"Moved {rf} -> results/")

    # 4. Generate combined results.xlsx inside results/
    print("Generating comprehensive combined results/results.xlsx...")
    combined_path = os.path.join('results', 'results.xlsx')
    
    wb = openpyxl.Workbook()
    # remove default sheet
    default_sheet = wb.active
    
    # Style definitions
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    sheets_data = [
        ('Obj1_Summary_Table3', os.path.join('results', 'objective_1_summary_table.csv'), 'Objective 1: Infeasible Rate Summary Across Conditions'),
        ('Obj1_All_Seeds_Detailed', os.path.join('results', 'objective_1_all_seeds_detailed.csv'), 'Objective 1: Detailed Infeasible Rates Across All 60 Seeds'),
        ('Obj2_Percolation_Summary', os.path.join('results', 'objective_2_percolation_summary.csv'), 'Objective 2: SFM Percolation Sweep Summary (18%-35% Penetration)'),
        ('Obj2_All_Seeds_Detailed', os.path.join('results', 'objective_2_percolation_all_seeds.csv'), 'Objective 2: Detailed Connectivity Across All 60 Runs')
    ]
    
    for sheet_name, csv_path, title_desc in sheets_data:
        ws = wb.create_sheet(title=sheet_name)
        if os.path.exists(csv_path):
            df = pd.read_csv(csv_path)
            
            # Write headers
            headers = list(df.columns)
            ws.append(headers)
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
            
            # Write data rows
            for r_idx, row in df.iterrows():
                row_vals = list(row)
                ws.append(row_vals)
                current_row = r_idx + 2
                for c_idx in range(1, len(row_vals) + 1):
                    cell = ws.cell(row=current_row, column=c_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    val = cell.value
                    if isinstance(val, (int, float)):
                        cell.alignment = center_align
                    else:
                        cell.alignment = center_align if any(k in str(val) for k in ['%', 'run', 'seed']) else left_align
            
            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        else:
            ws.append([f"CSV not found: {csv_path}"])

    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)
        
    wb.save(combined_path)
    print(f"Created combined Excel workbook: {combined_path}")

    # 5. Organize LaTeX files into latex/
    latex_files = [
        'report.tex',
        'PROJECT_REPORT.tex',
        'generate_report_pdf.py',
        'PROJECT_REPORT_NITK.pdf',
        'PROJECT_REPORT.pdf'
    ]
    for lf in latex_files:
        if os.path.exists(lf):
            dest = os.path.join('latex', lf)
            shutil.move(lf, dest)
            print(f"Moved {lf} -> latex/")

    # 6. Organize Docs into docs/
    doc_files = [
        'AboutProject.pdf',
        'IoV_Imputation_Implementation_Plan.pdf',
        'PROJECT_REPORT.md',
        'PPT_Information.md'
    ]
    for df in doc_files:
        if os.path.exists(df):
            dest = os.path.join('docs', df)
            shutil.move(df, dest)
            print(f"Moved {df} -> docs/")

    print("Project directory decluttering completed successfully!")

if __name__ == '__main__':
    main()
